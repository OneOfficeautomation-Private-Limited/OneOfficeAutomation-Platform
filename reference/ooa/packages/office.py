"""Office and integration packages: ooa.file, ooa.excel (basic, no Office needed),
ooa.rest. The Java equivalents use java.nio, Apache POI and the JDK HttpClient."""
from __future__ import annotations

import base64
import json
import re
import shutil
import urllib.error
import urllib.request
from decimal import Decimal
from pathlib import Path

from ..sdk import ActionException, Input, Output, action, condition, iterator, package, reveal
from ..sdk.types import make_table, to_jsonable


@package("ooa.file", "File", "1.0.0", description="Read, write, copy, move and delete files.")
class FilePackage:
    @action("read_text", "Read text", category="Files", description="Read a text file.",
            inputs=[Input("path", "File", description="File path."),
                    Input("encoding", "String", required=False, default="utf-8", description="Text encoding.")],
            outputs=[Output("text", "String", "File content.")], idempotent=True, capabilities=("filesystem",))
    def read_text(self, ctx, path, encoding="utf-8"):
        p = ctx.resolve_path(path)
        if not p.is_file():
            raise ActionException(f"File not found: {p}", code="FILE.NOT_FOUND", suggested_fix="Check the path and permissions.")
        return {"text": p.read_text(encoding=encoding)}

    @action("write_text", "Write text", category="Files", description="Write text to a file, overwriting or appending.",
            inputs=[Input("path", "File", description="File path."), Input("text", "String", description="Text to write."),
                    Input("append", "Boolean", required=False, default=False, description="Append instead of overwrite.")],
            capabilities=("filesystem",))
    def write_text(self, ctx, path, text, append=False):
        p = ctx.resolve_path(path)
        p.parent.mkdir(parents=True, exist_ok=True)
        with open(p, "a" if append else "w", encoding="utf-8") as fh:
            fh.write(text)
        return {}

    @action("copy", "Copy", category="Files", description="Copy a file.",
            inputs=[Input("source", "File", description="Source path."), Input("target", "File", description="Target path."),
                    Input("overwrite", "Boolean", required=False, default=False, description="Replace an existing target.")],
            capabilities=("filesystem",))
    def copy(self, ctx, source, target, overwrite=False):
        s, t = ctx.resolve_path(source), ctx.resolve_path(target)
        if t.exists() and not overwrite:
            raise ActionException(f"Target exists: {t}", code="FILE.EXISTS", suggested_fix="Set overwrite to true.")
        t.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(s, t)
        return {}

    @action("delete", "Delete", category="Files", description="Delete a file if it exists.",
            inputs=[Input("path", "File", description="File path.")], idempotent=True, capabilities=("filesystem",))
    def delete(self, ctx, path):
        ctx.resolve_path(path).unlink(missing_ok=True)
        return {}

    @action("list", "List files", category="Folders", description="List files in a folder matching a glob pattern.",
            inputs=[Input("folder", "File", description="Folder path."),
                    Input("pattern", "String", required=False, default="*", description="Glob pattern, e.g. *.pdf.")],
            outputs=[Output("files", "List", "Sorted list of file paths.")], idempotent=True, capabilities=("filesystem",))
    def list(self, ctx, folder, pattern="*"):
        return {"files": sorted(str(p) for p in ctx.resolve_path(folder).glob(pattern) if p.is_file())}

    @condition("exists", "File exists", description="True when the file exists.",
               inputs=[Input("path", "File", description="File path.")])
    def exists(self, ctx, path):
        return ctx.resolve_path(path).is_file()

    @iterator("files_in_folder", "Each file in folder", description="Yield each file path in a folder.",
              inputs=[Input("folder", "File", description="Folder path."),
                      Input("pattern", "String", required=False, default="*", description="Glob pattern.")])
    def files_in_folder(self, ctx, folder, pattern="*"):
        for p in sorted(ctx.resolve_path(folder).glob(pattern)):
            if p.is_file():
                yield str(p)


class _Workbook:
    def __init__(self, path: Path, wb):
        self.path, self.wb = path, wb

    def close(self):
        self.wb.close()


def _cell_value(v):
    if isinstance(v, float):
        return Decimal(repr(v))
    if isinstance(v, int) and not isinstance(v, bool):
        return Decimal(v)
    return v


@package("ooa.excel", "Excel (no Office needed)", "1.0.0",
         description="Read and write .xlsx workbooks without Microsoft Office installed.")
class ExcelPackage:
    _CELL = r"^[A-Z]{1,3}[1-9][0-9]*$"

    @action("open", "Open workbook", category="Workbook", description="Open an .xlsx workbook (or create it) as a named session.",
            inputs=[Input("path", "File", description="Workbook path."),
                    Input("session", "String", required=False, default="excel", description="Session name used by later steps."),
                    Input("create", "Boolean", required=False, default=False, description="Create the file if it is missing.")],
            capabilities=("filesystem",))
    def open(self, ctx, path, session="excel", create=False):
        import openpyxl
        p = ctx.resolve_path(path)
        if p.exists():
            wb = openpyxl.load_workbook(p)
        elif create:
            wb = openpyxl.Workbook()
        else:
            raise ActionException(f"Workbook not found: {p}", code="EXCEL.NOT_FOUND", suggested_fix="Set create to true or fix the path.")
        ctx.sessions[session] = _Workbook(p, wb)
        return {}

    def _sheet(self, ctx, session, sheet):
        wb = ctx.session(session, _Workbook).wb
        if sheet is None:
            return wb.active
        if sheet not in wb.sheetnames:
            raise ActionException(f"Sheet '{sheet}' not found", code="EXCEL.NO_SHEET",
                                  suggested_fix=f"Sheets: {', '.join(wb.sheetnames)}.")
        return wb[sheet]

    @action("get_cell", "Get cell", category="Cells", description="Read one cell value.",
            inputs=[Input("cell", "String", pattern=_CELL, example="B2", description="Cell address."),
                    Input("session", "String", required=False, default="excel", description="Workbook session."),
                    Input("sheet", "String", required=False, description="Sheet name; active sheet when empty.")],
            outputs=[Output("value", "Any", "Cell value.")], idempotent=True)
    def get_cell(self, ctx, cell, session="excel", sheet=None):
        return {"value": _cell_value(self._sheet(ctx, session, sheet)[cell].value)}

    @action("set_cell", "Set cell", category="Cells", description="Write one cell value.",
            inputs=[Input("cell", "String", pattern=_CELL, example="B2", description="Cell address."),
                    Input("value", "Any", description="Value to write."),
                    Input("session", "String", required=False, default="excel", description="Workbook session."),
                    Input("sheet", "String", required=False, description="Sheet name.")])
    def set_cell(self, ctx, cell, value, session="excel", sheet=None):
        v = to_jsonable(value)
        self._sheet(ctx, session, sheet)[cell] = v
        return {}

    @action("get_all_cells", "Get all cells", category="Cells", description="Read the used range as a Table; first row as headers.",
            inputs=[Input("session", "String", required=False, default="excel", description="Workbook session."),
                    Input("sheet", "String", required=False, description="Sheet name."),
                    Input("header", "Boolean", required=False, default=True, description="Treat first row as column names.")],
            outputs=[Output("table", "Table", "Sheet content.")], idempotent=True)
    def get_all_cells(self, ctx, session="excel", sheet=None, header=True):
        rows = [[_cell_value(c) for c in r] for r in self._sheet(ctx, session, sheet).iter_rows(values_only=True)]
        rows = [r for r in rows if any(c is not None for c in r)]
        if not rows:
            return {"table": make_table([], [])}
        cols = [str(c) for c in rows[0]] if header else [f"Column{i + 1}" for i in range(len(rows[0]))]
        return {"table": make_table(cols, rows[1:] if header else rows)}

    @action("write_table", "Write from Data Table", category="Cells", description="Write a Table starting at a cell, with headers.",
            inputs=[Input("table", "Table", description="Table to write."),
                    Input("start_cell", "String", required=False, default="A1", pattern=_CELL, description="Top-left cell."),
                    Input("session", "String", required=False, default="excel", description="Workbook session."),
                    Input("sheet", "String", required=False, description="Sheet name.")])
    def write_table(self, ctx, table, start_cell="A1", session="excel", sheet=None):
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        ws = self._sheet(ctx, session, sheet)
        col_l, row0 = coordinate_from_string(start_cell)
        col0 = column_index_from_string(col_l)
        for r, values in enumerate([table["columns"]] + table["rows"]):
            for c, v in enumerate(values):
                ws.cell(row=row0 + r, column=col0 + c, value=to_jsonable(v))
        return {}

    @action("add_sheet", "Create worksheet", category="Sheets", description="Create a worksheet, or clear it if it already exists.",
            inputs=[Input("name", "String", description="Worksheet name."),
                    Input("session", "String", required=False, default="excel", description="Workbook session."),
                    Input("replace", "Boolean", required=False, default=True, description="Clear an existing sheet with this name.")])
    def add_sheet(self, ctx, name, session="excel", replace=True):
        wb = ctx.session(session, _Workbook).wb
        if name in wb.sheetnames:
            if not replace:
                raise ActionException(f"Sheet '{name}' already exists", code="EXCEL.SHEET_EXISTS",
                                      suggested_fix="Set replace to true or choose another name.")
            del wb[name]
        wb.create_sheet(name)
        return {}

    @action("save", "Save", category="Workbook", description="Save the workbook, optionally to a new path.",
            inputs=[Input("session", "String", required=False, default="excel", description="Workbook session."),
                    Input("path", "File", required=False, description="Save-as path.")], capabilities=("filesystem",))
    def save(self, ctx, session="excel", path=None):
        wb = ctx.session(session, _Workbook)
        target = ctx.resolve_path(path) if path else wb.path
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.wb.save(target)
        return {}

    @action("close", "Close", category="Workbook", description="Close the workbook session without saving.",
            inputs=[Input("session", "String", required=False, default="excel", description="Workbook session.")])
    def close(self, ctx, session="excel"):
        wb = ctx.sessions.pop(session, None)
        if wb:
            wb.close()
        return {}


@package("ooa.rest", "REST Web Service", "1.0.0", description="Call HTTP/REST APIs with JSON, headers and authentication.")
class RestPackage:
    @action("request", "Send request", category="HTTP", description="Send an HTTP request and return status, headers and body.",
            inputs=[Input("url", "String", pattern=r"^https?://", description="Absolute URL."),
                    Input("method", "String", required=False, default="GET", enum=("GET", "POST", "PUT", "PATCH", "DELETE"), description="HTTP method."),
                    Input("headers", "Dictionary", required=False, description="Request headers."),
                    Input("json", "Any", required=False, description="JSON body (sets Content-Type)."),
                    Input("body", "String", required=False, description="Raw text body."),
                    Input("basic_user", "String", required=False, description="Basic auth user name."),
                    Input("basic_password", "Credential", required=False, description="Basic auth password from the vault."),
                    Input("bearer_token", "Credential", required=False, description="Bearer token from the vault."),
                    Input("fail_on_error", "Boolean", required=False, default=True, description="Raise REST.HTTP_ERROR on 4xx/5xx.")],
            outputs=[Output("status", "Number", "HTTP status code."), Output("body", "String", "Response text."),
                     Output("json", "Any", "Parsed JSON body or empty."), Output("headers", "Dictionary", "Response headers.")],
            timeout_seconds=60, capabilities=("network",))
    def request(self, ctx, url, method="GET", headers=None, json=None, body=None, basic_user=None,
                basic_password=None, bearer_token=None, fail_on_error=True):
        import json as _json
        hdrs = dict(headers or {})
        data = None
        if json is not None:
            data = _json.dumps(to_jsonable(json)).encode()
            hdrs.setdefault("Content-Type", "application/json")
        elif body is not None:
            data = body.encode()
        if basic_user and basic_password is not None:
            token = base64.b64encode(f"{basic_user}:{reveal(basic_password)}".encode()).decode()
            hdrs["Authorization"] = f"Basic {token}"
        if bearer_token is not None:
            hdrs["Authorization"] = f"Bearer {reveal(bearer_token)}"
        req = urllib.request.Request(url, data=data, method=method, headers=hdrs)
        safe_url = re.sub(r"([?&][^=]*(token|key|secret)[^=]*=)[^&]+", r"\1****", url, flags=re.I)
        ctx.log.info(f"{method} {safe_url}")
        try:
            with urllib.request.urlopen(req, timeout=55) as resp:
                status, text, rh = resp.status, resp.read().decode("utf-8", "replace"), dict(resp.headers)
        except urllib.error.HTTPError as e:
            status, text, rh = e.code, e.read().decode("utf-8", "replace"), dict(e.headers or {})
        except urllib.error.URLError as e:
            raise ActionException(f"Cannot reach {url}: {e.reason}", code="REST.UNREACHABLE", retryable=True,
                                  suggested_fix="Check the URL, proxy settings and that the device has network access.")
        try:
            parsed = _json.loads(text, parse_float=Decimal) if text else None
        except ValueError:
            parsed = None
        if fail_on_error and status >= 400:
            raise ActionException(f"{method} {url} returned HTTP {status}", code="REST.HTTP_ERROR",
                                  retryable=status in (429, 502, 503, 504), details={"status": status, "body": text[:500]})
        return {"status": status, "body": text, "json": parsed, "headers": rh}
